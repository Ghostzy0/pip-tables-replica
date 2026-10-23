"""Local, deliberately bounded data engine for Pip Tables.

Only files explicitly passed to ``import_file`` enter a project. User SQL runs
on a separate read-only DuckDB connection with external access disabled.
"""

from __future__ import annotations

import csv
import datetime as dt
import decimal
import hashlib
import json
import math
import os
import re
import tempfile
import threading
import time
import zipfile
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import duckdb
from defusedxml import ElementTree as safe_xml
from openpyxl import Workbook, load_workbook


_VERSION = 2
_MAX_INPUT_BYTES = 2 * 1024**3
_MAX_XLSX_ROWS = 1_000_000
_MAX_IMPORT_ROWS = 1_000_000
_IMPORT_TIMEOUT = 120
_MAX_XLSX_EXPANDED = 256 * 1024**2
_MAX_XLSX_SHARED_STRINGS = 16 * 1024**2
_MAX_XLSX_SHEET_XML = 128 * 1024**2
_MAX_XLSX_COLUMNS = 1024
_MAX_CELL_CHARS = 1_000_000
_MAX_PREVIEW = 10_000
_MAX_PREVIEW_BYTES = 8 * 1024**2
_MAX_PREVIEW_CELL_BYTES = 1024**2
_MAX_RESULT_COLUMNS = 512
_MAX_EXPORT_ROWS = 1_000_000
_EXPORT_TIMEOUT = 120
_NAME = re.compile(r"[^a-zA-Z0-9_]+")
_FORMULA = re.compile(r"^[\s\x00-\x1f]*[=+@-]")


def _ident(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def _name(value: str, fallback: str = "column") -> str:
    value = _NAME.sub("_", value.strip()).strip("_").lower()[:60]
    return value or fallback


def _unique(values: list[str], fallback: str = "column") -> list[str]:
    used: set[str] = set()
    result: list[str] = []
    for raw in values:
        base = _name(str(raw or ""), fallback)
        item = base
        index = 2
        while item.lower() in used:
            item = f"{base}_{index}"
            index += 1
        used.add(item.lower())
        result.append(item)
    return result


def _export_columns(values: list[str]) -> list[str]:
    """Preserve result labels while making duplicate keys and fields lossless."""
    used: set[str] = set()
    result = []
    for index, raw in enumerate(values, 1):
        base = str(raw) or f"column_{index}"
        candidate = base
        suffix = 2
        while candidate.lower() in used:
            candidate = f"{base}_{suffix}"
            suffix += 1
        used.add(candidate.lower())
        result.append(candidate)
    return result


def _safe_cell(value: Any) -> Any:
    """Stop spreadsheet programs from evaluating text supplied by a dataset."""
    if isinstance(value, str) and _FORMULA.match(value):
        return "'" + value
    return value


def _json_value(value: Any) -> Any:
    if isinstance(value, uuid.UUID):
        return str(value)
    if isinstance(value, decimal.Decimal):
        return str(value)
    if isinstance(value, (dt.date, dt.datetime, dt.time)):
        return value.isoformat()
    if isinstance(value, (bytes, bytearray, memoryview)):
        return bytes(value).hex()
    if isinstance(value, (dt.timedelta,)):
        return value.total_seconds()
    if isinstance(value, (float,)) and not math.isfinite(value):
        return None
    if isinstance(value, (list, tuple)):
        return [_json_value(v) for v in value]
    if isinstance(value, dict):
        return {str(k): _json_value(v) for k, v in value.items()}
    return value


def _xlsx_value(value: Any) -> Any:
    # Excel cells have narrower value/precision limits than DuckDB. Preserve
    # unsupported values as explicit text; never silently truncate a cell.
    if isinstance(value, (list, tuple, dict)):
        value = json.dumps(_json_value(value), ensure_ascii=False, allow_nan=False)
    elif isinstance(value, (uuid.UUID, dt.timedelta, bytes, bytearray, memoryview)):
        value = str(_json_value(value))
    elif isinstance(value, (dt.date, dt.datetime, dt.time)):
        value = value.isoformat()
    elif isinstance(value, decimal.Decimal):
        value = str(value)
    elif isinstance(value, int) and abs(value) > 2**53:
        value = str(value)
    elif isinstance(value, float) and not math.isfinite(value):
        value = str(value)
    if isinstance(value, str) and len(value) > 32766:
        raise ValueError("Excel cells must be under 32,767 characters; export JSON or Parquet instead")
    return _safe_cell(value)


@dataclass(frozen=True)
class QueryResult:
    columns: list[str]
    types: list[str]
    rows: list[tuple]
    total: int
    truncated: bool
    elapsed: float
    offset: int = 0


class Project:
    """An ordinary folder containing data.duckdb and atomic project.json."""

    def __init__(self, path: str | Path):
        self.path = Path(path).expanduser().resolve()
        self._db = self.path / "data.duckdb"
        self._meta_path = self.path / "project.json"
        if not self.path.is_dir() or not self._db.is_file() or not self._meta_path.is_file():
            raise FileNotFoundError(f"Not a Pip Tables project: {self.path}")
        if self._meta_path.stat().st_size > 2 * 1024**2:
            raise ValueError("Project metadata exceeds 2 MB")
        try:
            meta = json.loads(self._meta_path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise ValueError("Project metadata is unreadable") from exc
        if (not isinstance(meta, dict) or meta.get("version") != _VERSION
                or not isinstance(meta.get("draft"), str)
                or not isinstance(meta.get("history"), list)):
            raise ValueError("Unsupported or corrupt project metadata")
        if len(meta["draft"]) > 1_000_000 or len(meta["history"]) > 100:
            raise ValueError("Project metadata exceeds size limits")
        for entry in meta["history"]:
            if (not isinstance(entry, dict) or not isinstance(entry.get("sql"), str)
                    or len(entry["sql"]) > 1_000_000 or not isinstance(entry.get("rows"), int)
                    or entry["rows"] < 0 or not isinstance(entry.get("time"), str)):
                raise ValueError("Project query history is corrupt")
        try:
            with self._connect(read_only=True) as conn:
                schema = [r[0] for r in conn.execute("DESCRIBE _pip_sources").fetchall()]
                if schema != ["name", "source", "imported_at", "source_path"]:
                    raise ValueError("Unsupported project database schema; keep the folder and reopen with its original version")
                source_names = {row[0] for row in conn.execute("SELECT name FROM _pip_sources").fetchall()}
                objects = conn.execute("SELECT table_name, table_type FROM information_schema.tables WHERE table_schema='main'").fetchall()
                if (any(kind != "BASE TABLE" for _, kind in objects) or
                        {name for name, _ in objects} != source_names | {"_pip_sources"}):
                    raise ValueError("Project contains unrecognised database objects")
        except duckdb.Error as exc:
            raise ValueError("Project database is unreadable") from exc
        self._meta = meta
        self._lock = threading.RLock()
        self._active_lock = threading.Lock()
        self._active: duckdb.DuckDBPyConnection | None = None
        self._active_stop: threading.Event | None = None
        self._cancelled = threading.Event()
        self._preparing = False
        self._closed = False

    @classmethod
    def create(cls, path: str | Path) -> Project:
        root = Path(path).expanduser().resolve()
        if root.exists() and any(root.iterdir()):
            raise FileExistsError(f"Project folder is not empty: {root}")
        root.mkdir(parents=True, exist_ok=True)
        database = root / "data.duckdb"
        meta_path = root / "project.json"
        try:
            conn = duckdb.connect(str(database), config={"autoload_known_extensions": "false", "autoinstall_known_extensions": "false"})
            try:
                conn.execute("CREATE TABLE _pip_sources (name VARCHAR PRIMARY KEY, source VARCHAR NOT NULL, imported_at TIMESTAMP NOT NULL, source_path VARCHAR NOT NULL)")
            finally:
                conn.close()
            cls._atomic_json(meta_path, {"version": _VERSION, "draft": "", "history": []})
            return cls(root)
        except Exception:
            # Creation is the sole operation allowed to clean up these known new files.
            for suffix in ("data.duckdb", "data.duckdb.wal", "project.json"):
                (root / suffix).unlink(missing_ok=True)
            raise

    @staticmethod
    def _atomic_json(path: Path, data: dict) -> None:
        fd, temporary = tempfile.mkstemp(prefix=".project-", dir=path.parent)
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as stream:
                json.dump(data, stream, ensure_ascii=False, separators=(",", ":"))
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, path)
            directory = os.open(path.parent, os.O_RDONLY)
            try:
                os.fsync(directory)
            finally:
                os.close(directory)
        finally:
            Path(temporary).unlink(missing_ok=True)

    def _check_open(self) -> None:
        if self._closed:
            raise RuntimeError("Project is closed")

    def _connect(self, read_only: bool) -> duckdb.DuckDBPyConnection:
        config = {"autoload_known_extensions": "false", "autoinstall_known_extensions": "false",
                  "allow_unsigned_extensions": "false"}
        if read_only:
            config["enable_external_access"] = "false"
        conn = duckdb.connect(str(self._db), read_only=read_only, config=config)
        self._apply_limits(conn)
        return conn

    @staticmethod
    def _apply_limits(conn: duckdb.DuckDBPyConnection) -> None:
        conn.execute("SET memory_limit='512MB'")
        conn.execute("SET max_temp_directory_size='0B'")
        conn.execute("SET threads=2")

    def _write_meta(self, meta: dict) -> None:
        self._atomic_json(self._meta_path, meta)
        self._meta = meta

    @property
    def draft(self) -> str:
        with self._lock:
            self._check_open()
            return self._meta["draft"]

    def save_draft(self, sql: str) -> None:
        if not isinstance(sql, str) or len(sql) > 1_000_000:
            raise ValueError("Draft must be text under 1 MB")
        with self._lock:
            self._check_open()
            self._write_meta({**self._meta, "draft": sql})

    def history(self) -> list[dict]:
        with self._lock:
            self._check_open()
            return [dict(item) for item in self._meta["history"]]

    def tables(self) -> list[dict]:
        with self._lock:
            self._check_open()
            with self._connect(read_only=True) as conn:
                sources = conn.execute("SELECT name, source FROM _pip_sources ORDER BY imported_at, name").fetchall()
                result = []
                for name, source in sources:
                    count = conn.execute(f"SELECT count(*) FROM {_ident(name)}").fetchone()[0]
                    cols = conn.execute("SELECT column_name, data_type FROM information_schema.columns WHERE table_schema='main' AND table_name=? ORDER BY ordinal_position", [name]).fetchall()
                    result.append({"name": name, "rows": count,
                                   "columns": [{"name": n, "type": t} for n, t in cols], "source": source})
                return result

    def _existing_table(self, conn: duckdb.DuckDBPyConnection, table: str) -> None:
        if not isinstance(table, str) or conn.execute("SELECT count(*) FROM _pip_sources WHERE name=?", [table]).fetchone()[0] != 1:
            raise ValueError("Unknown imported table")

    def profile(self, table: str) -> list[dict]:
        with self._lock:
            self._check_open()
            with self._connect(read_only=True) as conn:
                self._existing_table(conn, table)
                expired = threading.Event()
                self._cancelled.clear()
                def stop():
                    expired.set()
                    conn.interrupt()
                timer = threading.Timer(10, stop)
                self._start_active(conn)
                timer.start()
                try:
                    cols = conn.execute("SELECT column_name, data_type FROM information_schema.columns WHERE table_schema='main' AND table_name=? ORDER BY ordinal_position", [table]).fetchall()
                    result = []
                    for name, kind in cols:
                        ident = _ident(name)
                        if self._cancelled.is_set():
                            raise InterruptedError("Profile stopped")
                        row = conn.execute(f"SELECT count(*) FILTER (WHERE {ident} IS NULL), count(DISTINCT {ident}), min({ident}), max({ident}) FROM {_ident(table)}").fetchone()
                        result.append({"name": name, "type": kind, "nulls": row[0], "distinct": row[1], "min": row[2], "max": row[3]})
                    return result
                except duckdb.Error as exc:
                    if expired.is_set():
                        raise TimeoutError("Profile timed out; use a focused SQL summary") from exc
                    if self._cancelled.is_set():
                        raise InterruptedError("Profile stopped") from exc
                    raise
                finally:
                    timer.cancel()
                    self._end_active(conn)

    @staticmethod
    def _validate_sql(conn: duckdb.DuckDBPyConnection, sql: str) -> str:
        if not isinstance(sql, str) or not sql.strip() or len(sql) > 1_000_000:
            raise ValueError("Enter one read-only SQL query")
        try:
            statements = conn.extract_statements(sql)
        except duckdb.Error as exc:
            raise ValueError(f"Invalid SQL: {exc}") from exc
        normalized = sql.strip().rstrip(";")
        if len(statements) == 1 and statements[0].type == duckdb.StatementType.SELECT:
            return normalized
        # DuckDB reports standalone PIVOT as CREATE because it rewrites the
        # result internally. Exactly two internal statements are expected.
        if (len(statements) == 2 and
                [s.type for s in statements] == [duckdb.StatementType.CREATE, duckdb.StatementType.SELECT]
                and re.match(r"(?is)^PIVOT\b", normalized)):
            wrapped = f"SELECT * FROM ({normalized}) AS _pip_pivot"
            parsed = conn.extract_statements(wrapped)
            if (len(parsed) == 2 and
                    [s.type for s in parsed] == [duckdb.StatementType.CREATE, duckdb.StatementType.SELECT]):
                return wrapped
        raise ValueError("Only one SELECT or PIVOT query is allowed")

    def _start_active(self, conn: duckdb.DuckDBPyConnection) -> None:
        with self._active_lock:
            if self._active_stop:
                self._active_stop.set()
            finished = threading.Event()
            self._active_stop = finished
            self._active = conn
        def watch_cancel():
            # DuckDB clears an interrupt between statements. Retain a Stop
            # request across count/preview/staging transitions, including the
            # narrow gap immediately before execute starts.
            while not finished.wait(.02):
                if self._cancelled.is_set():
                    with self._active_lock:
                        if self._active is conn:
                            conn.interrupt()
        threading.Thread(target=watch_cancel, daemon=True).start()

    def _end_active(self, conn: duckdb.DuckDBPyConnection) -> None:
        with self._active_lock:
            if self._active is conn:
                if self._active_stop:
                    self._active_stop.set()
                self._active = None

    def cancel(self) -> None:
        """Interrupt the current import, query or export from another thread."""
        with self._active_lock:
            if self._active is not None or self._preparing:
                self._cancelled.set()
            if self._active is not None:
                self._active.interrupt()

    def _set_preparing(self, value: bool) -> None:
        with self._active_lock:
            self._preparing = value

    @staticmethod
    def _source(path: str | Path, sheet: str | None, encoding: str) -> tuple[Path, str]:
        source = Path(path).expanduser().resolve(strict=True)
        if not source.is_file() or source.stat().st_size > _MAX_INPUT_BYTES:
            raise ValueError("Select a regular file under 2 GB")
        if any(character in str(source) for character in "*?[]"):
            raise ValueError("File names containing glob characters are unsupported")
        kind = source.suffix.lower()
        if kind not in {".csv", ".xlsx", ".json", ".jsonl", ".ndjson", ".parquet"}:
            raise ValueError("Supported imports: CSV, XLSX, JSON, NDJSON, Parquet")
        if sheet is not None and kind != ".xlsx":
            raise ValueError("A sheet can only be selected for XLSX")
        if not isinstance(encoding, str) or not encoding or "\x00" in encoding:
            raise ValueError("Invalid encoding")
        if kind == ".xlsx":
            # Central-directory checks happen before any worksheet XML is read.
            with zipfile.ZipFile(source) as archive:
                entries = archive.infolist()
                if (len(entries) > 10_000 or sum(item.file_size for item in entries) > _MAX_XLSX_EXPANDED
                        or any(item.filename == "xl/workbook.xml" and item.file_size > 10 * 1024**2 for item in entries)):
                    raise ValueError("XLSX expanded size exceeds safety limit")
                if any(item.filename == "xl/sharedStrings.xml" and item.file_size > _MAX_XLSX_SHARED_STRINGS
                       for item in entries):
                    raise ValueError("XLSX shared strings exceed 16 MB safety limit")
                if any(item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml")
                       and item.file_size > _MAX_XLSX_SHEET_XML for item in entries):
                    raise ValueError("XLSX worksheet XML exceeds 128 MB safety limit")
        return source, kind

    def _fingerprint(self, source: Path) -> str:
        if self._cancelled.is_set():
            raise InterruptedError("File inspection stopped")
        digest = hashlib.sha256()
        with source.open("rb") as stream:
            while block := stream.read(1024 * 1024):
                if self._cancelled.is_set():
                    raise InterruptedError("File inspection stopped")
                digest.update(block)
        if self._cancelled.is_set():
            raise InterruptedError("File inspection stopped")
        return "sha256:" + digest.hexdigest()

    @staticmethod
    def _sheet_names(source: Path) -> list[str]:
        # Workbook metadata is independent of sharedStrings and worksheet
        # cells, both of which can be expensive to parse before choosing a
        # sheet. Bound actual decompressed bytes, not just ZIP declarations.
        with zipfile.ZipFile(source) as archive:
            with archive.open("xl/workbook.xml") as stream:
                xml = stream.read(10 * 1024**2 + 1)
        if len(xml) > 10 * 1024**2:
            raise ValueError("XLSX workbook metadata exceeds 10 MB")
        root = safe_xml.fromstring(xml)
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        return [item.attrib["name"] for item in root.findall(f"{ns}sheets/{ns}sheet") if "name" in item.attrib]

    def worksheet_names(self, path: str | Path) -> list[str]:
        """Read XLSX workbook metadata; worksheet cells are not parsed."""
        with self._lock:
            self._check_open()
            self._cancelled.clear()
            self._set_preparing(True)
            try:
                source, kind = self._source(path, None, "utf-8")
                if kind != ".xlsx":
                    raise ValueError("Choose an XLSX workbook")
                names = self._sheet_names(source)
                if self._cancelled.is_set():
                    raise InterruptedError("Workbook inspection stopped")
                return names
            finally:
                self._set_preparing(False)

    def query(self, sql: str, limit: int = 1000, timeout: float = 10,
              offset: int = 0, record_history: bool = True) -> QueryResult:
        if not isinstance(limit, int) or not 1 <= limit <= _MAX_PREVIEW:
            raise ValueError(f"Preview limit must be 1–{_MAX_PREVIEW}")
        if not isinstance(timeout, (int, float)) or not 0 < timeout <= 120:
            raise ValueError("Timeout must be 0–120 seconds")
        if isinstance(offset, bool) or not isinstance(offset, int) or not 0 <= offset <= 2**63 - 1:
            raise ValueError("Offset must be a nonnegative integer")
        if not isinstance(record_history, bool):
            raise ValueError("record_history must be a boolean")
        with self._lock:
            self._check_open()
            start = time.monotonic()
            with self._connect(read_only=True) as conn:
                normalized = self._validate_sql(conn, sql)
                expired = threading.Event()
                self._cancelled.clear()

                def stop() -> None:
                    expired.set()
                    conn.interrupt()

                timer = threading.Timer(timeout, stop)
                self._start_active(conn)
                timer.start()
                try:
                    # The exact count and preview share one deadline. No full result is
                    # materialised in Python; DuckDB also has a memory/spill cap.
                    total = conn.execute(f"SELECT count(*) FROM ({normalized}) AS _pip_count").fetchone()[0]
                    # Bound the native-engine snapshot before copying any result
                    # values into Python/Qt. Temp tables are allowed on a read-only
                    # connection; the persistent project remains read-only.
                    conn.execute(f"CREATE TEMP TABLE _pip_preview_data AS SELECT * FROM ({normalized}) AS _pip_preview LIMIT ? OFFSET ?", [limit, offset])
                    cursor = conn.execute("SELECT * FROM _pip_preview_data LIMIT 0")
                    columns = [item[0] for item in cursor.description]
                    types = [str(item[1]) for item in cursor.description]
                    if len(columns) > _MAX_RESULT_COLUMNS:
                        raise ValueError("Choose at most 512 result columns")
                    sizes = [f"COALESCE(octet_length(encode(CAST({_ident(col)} AS VARCHAR))), 0)" for col in columns]
                    if sizes:
                        size, cell = conn.execute(f"SELECT COALESCE(SUM({' + '.join(sizes)}),0), COALESCE(MAX(GREATEST({','.join(sizes)})),0) FROM _pip_preview_data").fetchone()
                        if size > _MAX_PREVIEW_BYTES or cell > _MAX_PREVIEW_CELL_BYTES:
                            raise ValueError("Preview exceeds 8 MB or a cell exceeds 1 MB. Select fewer columns, shorten text or add LIMIT")
                    rows = conn.execute("SELECT * FROM _pip_preview_data").fetchmany(limit)
                    if expired.is_set():
                        raise TimeoutError("Query timed out")
                    if self._cancelled.is_set():
                        raise InterruptedError("Query stopped")
                except duckdb.Error as exc:
                    if expired.is_set():
                        raise TimeoutError("Query timed out") from exc
                    if self._cancelled.is_set():
                        raise InterruptedError("Query stopped") from exc
                    raise ValueError(f"Query failed: {exc}") from exc
                finally:
                    timer.cancel()
                    self._end_active(conn)
            elapsed = time.monotonic() - start
            if record_history:
                entry = {"sql": normalized, "rows": total, "time": dt.datetime.now(dt.timezone.utc).isoformat()}
                self._write_meta({**self._meta, "history": ([entry] + self._meta["history"])[:100]})
            return QueryResult(columns, types, rows, total, total > offset + len(rows), elapsed, offset)

    @staticmethod
    def _load_file(conn: duckdb.DuckDBPyConnection, source: Path, kind: str,
                   name: str, sheet: str | None, encoding: str, all_text: bool,
                   cancelled: threading.Event) -> None:
        # Both inspection and committed import use this exact loader.
        if kind == ".csv":
            conn.execute(f"CREATE TABLE {_ident(name)} AS SELECT * FROM read_csv(?, header=true, auto_detect=true, all_varchar=?, encoding=?) LIMIT ?",
                         [str(source), bool(all_text), encoding, _MAX_IMPORT_ROWS + 1])
        elif kind in {".json", ".jsonl", ".ndjson"}:
            conn.execute(f"CREATE TABLE {_ident(name)} AS SELECT * FROM read_json_auto(?, maximum_object_size=10485760) LIMIT ?", [str(source), _MAX_IMPORT_ROWS + 1])
        elif kind == ".parquet":
            conn.execute(f"CREATE TABLE {_ident(name)} AS SELECT * FROM read_parquet(?) LIMIT ?", [str(source), _MAX_IMPORT_ROWS + 1])
        else:
            Project._import_xlsx(conn, source, name, sheet, all_text, cancelled)

    def inspect_file(self, path: str | Path, sheet: str | None = None,
                     encoding: str = "utf-8", all_text: bool = False) -> dict:
        """Inspect selected input in an isolated temporary DB without saving it."""
        with self._lock:
            self._check_open()
            self._cancelled.clear()
            self._set_preparing(True)
            expired = threading.Event()

            def stop() -> None:
                expired.set()
                self.cancel()

            timer = threading.Timer(_IMPORT_TIMEOUT, stop)
            timer.start()
            try:
                source, kind = self._source(path, sheet, encoding)
                before = self._fingerprint(source)
                sheets: list[str] = []
                if kind == ".xlsx":
                    sheets = self._sheet_names(source)
                    if sheet is None and len(sheets) != 1:
                        if self._fingerprint(source) != before:
                            raise ValueError("Source changed during inspection")
                        if expired.is_set():
                            raise TimeoutError("File inspection timed out")
                        return {"columns": [], "rows": [], "total": 0,
                                "sheets": sheets, "fingerprint": before}
                with tempfile.TemporaryDirectory(prefix="pip-tables-inspect-") as folder:
                    stage = duckdb.connect(str(Path(folder) / "inspect.duckdb"), config={
                        "autoload_known_extensions": "false", "autoinstall_known_extensions": "false",
                        "allow_unsigned_extensions": "false"})
                    try:
                        self._apply_limits(stage)
                        self._start_active(stage)
                        try:
                            self._load_file(stage, source, kind, "_pip_inspect", sheet, encoding, all_text, self._cancelled)
                            if expired.is_set():
                                raise TimeoutError("File inspection timed out")
                            if self._cancelled.is_set():
                                raise InterruptedError("File inspection stopped")
                            columns = [{"name": n, "type": t} for n, t in stage.execute(
                                "SELECT column_name, data_type FROM information_schema.columns WHERE table_schema='main' AND table_name='_pip_inspect' ORDER BY ordinal_position").fetchall()]
                            if not columns or len(columns) > _MAX_RESULT_COLUMNS:
                                raise ValueError("Inspection needs 1–512 columns")
                            total = stage.execute("SELECT count(*) FROM _pip_inspect").fetchone()[0]
                            if total > _MAX_IMPORT_ROWS:
                                raise ValueError("Import row limit exceeded (1,000,000)")
                            names = [item["name"] for item in columns]
                            sizes = [f"COALESCE(octet_length(encode(CAST({_ident(name)} AS VARCHAR))), 0)" for name in names]
                            size, cell = stage.execute(f"SELECT COALESCE(SUM({' + '.join(sizes)}),0), COALESCE(MAX(GREATEST({','.join(sizes)})),0) FROM (SELECT * FROM _pip_inspect LIMIT 20)").fetchone()
                            if size > _MAX_PREVIEW_BYTES or cell > _MAX_PREVIEW_CELL_BYTES:
                                raise ValueError("Inspection preview exceeds 8 MB or a cell exceeds 1 MB")
                            rows = stage.execute("SELECT * FROM _pip_inspect LIMIT 20").fetchmany(20)
                        finally:
                            self._end_active(stage)
                    finally:
                        stage.close()
                if expired.is_set():
                    raise TimeoutError("File inspection timed out")
                if self._cancelled.is_set():
                    raise InterruptedError("File inspection stopped")
                if self._fingerprint(source) != before:
                    raise ValueError("Source changed during inspection")
                return {"columns": columns, "rows": rows, "total": total,
                        "sheets": sheets, "fingerprint": before}
            except duckdb.Error as exc:
                if expired.is_set():
                    raise TimeoutError("File inspection timed out") from exc
                if self._cancelled.is_set():
                    raise InterruptedError("File inspection stopped") from exc
                raise ValueError(f"Inspection failed: {exc}") from exc
            except InterruptedError as exc:
                if expired.is_set():
                    raise TimeoutError("File inspection timed out") from exc
                raise
            finally:
                timer.cancel()
                self._set_preparing(False)

    def import_file(self, path: str | Path, sheet: str | None = None,
                    encoding: str = "utf-8", all_text: bool = False,
                    expected_fingerprint: str | None = None) -> str:
        with self._lock:
            self._check_open()
            self._cancelled.clear()
            self._set_preparing(True)
            expired = threading.Event()

            def stop() -> None:
                expired.set()
                self.cancel()

            timer = threading.Timer(_IMPORT_TIMEOUT, stop)
            timer.start()
            try:
                source, kind = self._source(path, sheet, encoding)
                if expected_fingerprint is not None:
                    if (not isinstance(expected_fingerprint, str) or
                            not re.fullmatch(r"sha256:[0-9a-f]{64}", expected_fingerprint)):
                        raise ValueError("Invalid inspection fingerprint")
                    if self._fingerprint(source) != expected_fingerprint:
                        raise ValueError("Source changed since inspection; inspect it again")
                connection = self._connect(read_only=False)
            except BaseException:
                timer.cancel()
                self._set_preparing(False)
                if expired.is_set():
                    raise TimeoutError("Import timed out")
                raise
            with connection as conn:
                existing = {r[0].lower() for r in conn.execute("SELECT name FROM _pip_sources").fetchall()}
                base = _name(source.stem, "table")
                if base.startswith("_pip_"):
                    base = "data_" + base
                name = base
                counter = 2
                while name.lower() in existing:
                    name = f"{base}_{counter}"
                    counter += 1
                self._start_active(conn)
                try:
                    conn.execute("BEGIN")
                    self._load_file(conn, source, kind, name, sheet, encoding, all_text, self._cancelled)
                    if expired.is_set():
                        raise TimeoutError("Import timed out")
                    if self._cancelled.is_set():
                        raise InterruptedError("Import stopped")
                    count = conn.execute(f"SELECT count(*) FROM {_ident(name)}").fetchone()[0]
                    if count > _MAX_IMPORT_ROWS:
                        raise ValueError("Import row limit exceeded (1,000,000)")
                    cols = conn.execute("SELECT count(*) FROM information_schema.columns WHERE table_schema='main' AND table_name=?", [name]).fetchone()[0]
                    if cols == 0:
                        raise ValueError("Input has no columns")
                    if expected_fingerprint is not None and self._fingerprint(source) != expected_fingerprint:
                        raise ValueError("Source changed during import; inspect it again")
                    conn.execute("INSERT INTO _pip_sources VALUES (?, ?, current_timestamp, ?)", [name, source.name, str(source)])
                    conn.execute("COMMIT")
                    return name
                except Exception as exc:
                    try:
                        conn.execute("ROLLBACK")
                    except duckdb.Error:
                        pass
                    if self._cancelled.is_set():
                        if expired.is_set():
                            raise TimeoutError("Import timed out") from exc
                        raise InterruptedError("Import stopped") from exc
                    raise
                finally:
                    timer.cancel()
                    self._end_active(conn)
                    self._set_preparing(False)

    @staticmethod
    def _import_xlsx(conn: duckdb.DuckDBPyConnection, source: Path, name: str,
                     sheet: str | None, all_text: bool, cancelled: threading.Event) -> None:
        book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        try:
            if sheet is None:
                if len(book.sheetnames) != 1:
                    raise ValueError(f"Select an XLSX sheet: {', '.join(book.sheetnames)}")
                sheet = book.sheetnames[0]
            if sheet not in book.sheetnames:
                raise ValueError("XLSX sheet not found")
            page = book[sheet]
            if ((page.max_column or 0) > _MAX_XLSX_COLUMNS
                    or (page.max_row or 0) > _MAX_XLSX_ROWS + 1):
                raise ValueError("XLSX dimensions exceed import limit")
            iterator = page.iter_rows()
            header = next(iterator, None)
            if not header:
                raise ValueError("Empty XLSX sheet")
            if len(header) > _MAX_XLSX_COLUMNS:
                raise ValueError("XLSX column limit exceeded")
            columns = _unique([str(cell.value or "") for cell in header])
            if any(cell.data_type == "f" for cell in header):
                raise ValueError("XLSX formula headers are unsupported")
            kinds = ["NULL"] * len(columns)
            has_value = [False] * len(columns)
            wide_integer = [False] * len(columns)
            for row_number, row in enumerate(iterator, 1):
                if cancelled.is_set():
                    raise InterruptedError("Import stopped")
                if row_number > _MAX_XLSX_ROWS:
                    raise ValueError("XLSX row limit exceeded")
                for index, cell in enumerate(row[:len(columns)]):
                    if cell.data_type == "f":
                        raise ValueError("XLSX formulas are unsupported; save values before import")
                    value = cell.value
                    if value is None:
                        continue
                    if isinstance(value, str) and len(value) > _MAX_CELL_CHARS:
                        raise ValueError("XLSX cell exceeds 1 MB")
                    has_value[index] = True
                    incoming = ("VARCHAR" if all_text or isinstance(value, str) else
                                "BOOLEAN" if isinstance(value, bool) else
                                "BIGINT" if isinstance(value, int) else
                                "DOUBLE" if isinstance(value, float) else
                                "TIMESTAMP" if isinstance(value, dt.datetime) else
                                "DATE" if isinstance(value, dt.date) else "VARCHAR")
                    if isinstance(value, int) and not isinstance(value, bool):
                        wide_integer[index] |= abs(value) > 2**53
                        if not -(2**63) <= value < 2**63:
                            incoming = "VARCHAR"
                    current = kinds[index]
                    if current == "NULL":
                        kinds[index] = incoming
                    elif current == "VARCHAR" or incoming == "VARCHAR":
                        kinds[index] = "VARCHAR"
                    elif current == incoming:
                        pass
                    elif {current, incoming} <= {"BIGINT", "DOUBLE"}:
                        kinds[index] = "DOUBLE" if "DOUBLE" in {current, incoming} else "BIGINT"
                    elif {current, incoming} <= {"DATE", "TIMESTAMP"}:
                        kinds[index] = "TIMESTAMP" if "TIMESTAMP" in {current, incoming} else "DATE"
                    else:
                        kinds[index] = "VARCHAR"
            kinds = ["VARCHAR" if all_text or not has_value[i] else value for i, value in enumerate(kinds)]
            # A double cannot exactly represent every integer above 2**53.
            # Keep conflicting Excel numerics as text so the owner can cast
            # explicitly instead of silently losing significant digits.
            kinds = ["VARCHAR" if wide_integer[i] and kind == "DOUBLE" else kind for i, kind in enumerate(kinds)]
            definition = ", ".join(f"{_ident(col)} {kind}" for col, kind in zip(columns, kinds))
            conn.execute(f"CREATE TABLE {_ident(name)} ({definition})")
        finally:
            book.close()
        book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        try:
            page = book[sheet]
            iterator = page.iter_rows(values_only=True)
            next(iterator, None)
            placeholders = ",".join("?" for _ in columns)
            batch: list[tuple] = []
            for row in iterator:
                if cancelled.is_set():
                    raise InterruptedError("Import stopped")
                values = []
                for index, kind in enumerate(kinds):
                    value = row[index] if index < len(row) else None
                    values.append(str(value) if all_text and value is not None else value)
                batch.append(tuple(values))
                if len(batch) == 1000:
                    conn.executemany(f"INSERT INTO {_ident(name)} VALUES ({placeholders})", batch)
                    batch.clear()
            if batch:
                conn.executemany(f"INSERT INTO {_ident(name)} VALUES ({placeholders})", batch)
        finally:
            book.close()

    def export(self, sql: str, path: str | Path, format: str) -> int:
        """Write a complete result atomically. One million rows is the safety cap."""
        kind = format.lower().lstrip(".")
        if kind not in {"csv", "xlsx", "json", "parquet"}:
            raise ValueError("Export format must be CSV, XLSX, JSON or Parquet")
        target = Path(path).expanduser().resolve()
        if not target.parent.is_dir():
            raise FileNotFoundError(target.parent)
        with self._lock:
            self._check_open()
            # A chosen destination must never corrupt a project or any selected
            # input. Resolve symlinks and also compare inode aliases (hardlinks).
            if target.is_relative_to(self.path):
                raise ValueError("Export outside the project folder to keep the project safe")
            with self._connect(read_only=True) as guard:
                sources = [Path(row[0]) for row in guard.execute("SELECT source_path FROM _pip_sources").fetchall()]
            protected = sources + [self._db, self._meta_path, self.path / "data.duckdb.wal"]
            for source in protected:
                if target == source.resolve() or (target.exists() and source.exists() and os.path.samefile(target, source)):
                    raise ValueError("Choose a new destination; original inputs and project files are protected")
            if target.exists():
                raise FileExistsError("Exports never replace existing files. Choose a new filename")
            fd, temporary = tempfile.mkstemp(prefix=".pip-export-", suffix="." + kind, dir=target.parent)
            os.close(fd)
            writer_temporary = None
            try:
                # Parquet needs DuckDB's native writer for bounded, vectorized
                # output. Start with external access enabled only long enough to
                # grant that connection access to this one random temporary file,
                # then permanently disable external access before validating or
                # executing any user SQL.
                if kind == "parquet":
                    conn = duckdb.connect(
                        str(self._db), read_only=True,
                        config={"autoload_known_extensions": "false",
                                "autoinstall_known_extensions": "false",
                                "allow_unsigned_extensions": "false"},
                    )
                    try:
                        self._apply_limits(conn)
                        temporary_path = Path(temporary)
                        # DuckDB writes to a sibling tmp_<name> before its final
                        # atomic rename, so both exact generated filenames must be
                        # authorized. Neither their parent directory nor a wildcard
                        # is granted.
                        writer_temporary = str(temporary_path.with_name("tmp_" + temporary_path.name))
                        if Path(writer_temporary).exists():
                            writer_temporary = None
                            raise FileExistsError("Export temporary sibling already exists; choose another filename")
                        conn.execute("SET allowed_paths = ?", [[temporary, writer_temporary]])
                        conn.execute("SET enable_external_access = false")
                    except Exception:
                        conn.close()
                        raise
                else:
                    conn = self._connect(read_only=True)
                with conn:
                    normalized = self._validate_sql(conn, sql)
                    self._cancelled.clear()
                    self._start_active(conn)
                    expired = threading.Event()

                    def stop() -> None:
                        expired.set()
                        self.cancel()

                    timer = threading.Timer(_EXPORT_TIMEOUT, stop)
                    timer.start()
                    written = 0
                    def batches(cursor):
                        nonlocal written
                        while rows := cursor.fetchmany(1000):
                            written += len(rows)
                            if written > _MAX_EXPORT_ROWS:
                                raise ValueError("Export exceeds one million rows")
                            if expired.is_set():
                                raise TimeoutError("Export timed out")
                            if self._cancelled.is_set():
                                raise InterruptedError("Export stopped")
                            yield rows
                    try:
                        if kind == "parquet":
                            # The row cap is part of the one materializing query:
                            # COPY returns the number actually written, so volatile
                            # queries are never run once for a count and again for
                            # the export. DuckDB's vectorized COPY preserves types
                            # and nested values without Python row conversion.
                            bounded = f"SELECT * FROM ({normalized}) AS _pip_export LIMIT {_MAX_EXPORT_ROWS + 1}"
                            stats = conn.execute(
                                f"COPY ({bounded}) TO ? (FORMAT PARQUET, RETURN_STATS)",
                                [temporary],
                            ).fetchone()
                            written = int(stats[1])
                            if written > _MAX_EXPORT_ROWS:
                                raise ValueError(f"Export exceeds {_MAX_EXPORT_ROWS:,} rows")
                        else:
                            count = conn.execute(f"SELECT count(*) FROM ({normalized}) AS _pip_count").fetchone()[0]
                            if count > _MAX_EXPORT_ROWS:
                                raise ValueError(f"Export exceeds {_MAX_EXPORT_ROWS:,} rows")
                            cursor = conn.execute(f"SELECT * FROM ({normalized}) AS _pip_export")
                            columns = _export_columns([item[0] for item in cursor.description])
                            if kind == "csv":
                                with open(temporary, "w", newline="", encoding="utf-8") as stream:
                                    writer = csv.writer(stream)
                                    writer.writerow([_safe_cell(name) for name in columns])
                                    for rows in batches(cursor):
                                        if expired.is_set(): raise TimeoutError("Export timed out")
                                        if self._cancelled.is_set(): raise InterruptedError("Export stopped")
                                        writer.writerows([_safe_cell(v) for v in row] for row in rows)
                            elif kind == "json":
                                with open(temporary, "w", encoding="utf-8") as stream:
                                    stream.write("[\n")
                                    first = True
                                    for rows in batches(cursor):
                                        if expired.is_set(): raise TimeoutError("Export timed out")
                                        if self._cancelled.is_set(): raise InterruptedError("Export stopped")
                                        for row in rows:
                                            if not first: stream.write(",\n")
                                            json.dump(dict(zip(columns, map(_json_value, row))), stream, ensure_ascii=False, allow_nan=False)
                                            first = False
                                    stream.write("\n]\n")
                            else:
                                book = Workbook(write_only=True)
                                sheet = book.create_sheet("Data")
                                sheet.append([_xlsx_value(name) for name in columns])
                                for rows in batches(cursor):
                                    if expired.is_set(): raise TimeoutError("Export timed out")
                                    if self._cancelled.is_set(): raise InterruptedError("Export stopped")
                                    for row in rows:
                                        sheet.append([_xlsx_value(v) for v in row])
                                book.save(temporary)
                        if expired.is_set():
                            raise TimeoutError("Export timed out")
                        if self._cancelled.is_set():
                            raise InterruptedError("Export stopped")
                    except duckdb.Error as exc:
                        if expired.is_set():
                            raise TimeoutError("Export timed out") from exc
                        if self._cancelled.is_set():
                            raise InterruptedError("Export stopped") from exc
                        raise ValueError(f"Export failed: {exc}") from exc
                    finally:
                        timer.cancel()
                        self._end_active(conn)
                # Atomic create-if-absent also protects a destination created
                # by another process while export was in progress.
                with open(temporary, "rb") as completed:
                    os.fsync(completed.fileno())
                os.link(temporary, target)
                return written
            finally:
                Path(temporary).unlink(missing_ok=True)
                if writer_temporary is not None:
                    Path(writer_temporary).unlink(missing_ok=True)

    def close(self) -> None:
        with self._lock:
            self._closed = True
